from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# Function to create Resume
def create_resume():
    doc = Document()

    # Header
    title = doc.add_heading('CURRICULUM VITAE', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Personal Information
    doc.add_heading('PERSONAL DETAILS', level=1)
    p = doc.add_paragraph()
    p.add_run('Name: ').bold = True
    p.add_run('Dr. Rakesh Singh\n')
    p.add_run('e-HRMS / PIS Code: ').bold = True
    p.add_run('256035\n')
    p.add_run('Father\'s Name: ').bold = True
    p.add_run('Late S.B. Singh\n')
    p.add_run('Date of Birth: ').bold = True
    p.add_run('04 Nov 1960\n')
    p.add_run('Address: ').bold = True
    p.add_run('D-141 Rajajipuram, Lucknow - 226017\n')
    p.add_run('Mobile: ').bold = True
    p.add_run('9415136011\n')
    p.add_run('Email: ').bold = True
    p.add_run('drrsingh79@gmail.com\n')
    p.add_run('PAN Number: ').bold = True
    p.add_run('AWSPS6115Q\n')

    # Education
    doc.add_heading('EDUCATIONAL QUALIFICATION', level=1)
    doc.add_paragraph('Diploma in Orthopaedics (D. Ortho) - Kanpur University, 1990', style='List Bullet')
    doc.add_paragraph('MBBS - Kanpur University, 1987', style='List Bullet')
    doc.add_paragraph('B.Sc (Botany, Zoology) - Lucknow University, 1979', style='List Bullet')
    doc.add_paragraph('Intermediate (12th) - 1977', style='List Bullet')
    doc.add_paragraph('High School (10th) - 1975', style='List Bullet')

    # Professional Experience
    doc.add_heading('PROFESSIONAL EXPERIENCE & SERVICE HISTORY', level=1)

    # Extension
    p_ext = doc.add_paragraph()
    p_ext.add_run('Post-Retirement Extension (2022 – 2024): ').bold = True
    p_ext.add_run(
        'Served as Orthopaedic Surgeon for a 2-year extension at R.L.B. Hospital, Lucknow, post superannuation.')

    # Retirement
    p_ret = doc.add_paragraph()
    p_ret.add_run('Superannuation (30 Nov 2022): ').bold = True
    p_ret.add_run('Retired as Senior Consultant from Dr. SPM Hospital (Civil Hospital), Lucknow.')

    # Career Postings
    doc.add_paragraph('Senior Consultant - Dr. SPM Hospital, Lucknow (July 2015 – Nov 2022)', style='List Bullet')
    doc.add_paragraph('Senior Consultant - District Male Hospital, Hardoi (June 2013 – June 2015)', style='List Bullet')
    doc.add_paragraph(
        'Medical Officer - Govt. Combined Hospital Jaspur, Nainital (Regularized: 02-09-1992). GO No: 3834/CHIKI-4-92-1510/91 DATED 06-08-1992',
        style='List Bullet')
    doc.add_paragraph(
        'Joined PMHS on Adhoc Basis - Surgeon at Hasanganj CHO (22-09-1990). GO No: 7893 SEC-4/PAANCH-450/88 T.C. LKO DATED 07-09-90',
        style='List Bullet')

    doc.save('Resume_Dr_Rakesh_Singh.docx')


# Function to create Job Application
def create_application():
    doc = Document()

    # Date and To
    doc.add_paragraph('Date: 12 February 2026')
    doc.add_paragraph(
        'To,\nThe Chief Medical Officer (CMO),\nPandit Deendayal Upadhyay Bhawan,\n1, Chakbast Road, Kaiserbagh,\nLucknow - 226018')

    # Subject
    subj = doc.add_paragraph()
    subj.add_run(
        'Subject: Application for Empanelment as Specialist (Orthopaedic Surgeon) for Polyclinics.').bold = True

    # Content
    doc.add_paragraph('Respected Sir,')
    doc.add_paragraph(
        'In reference to the advertisement issued by your office for the empanelment of specialist doctors in 14 identified Polyclinics in Lucknow, I wish to offer my services as an Orthopaedic Surgeon.')

    doc.add_paragraph(
        'I am a retired Senior Consultant from the PMHS cadre (e-HRMS: 256035) with extensive experience in Orthopaedic surgery. After my retirement from Dr. SPM Hospital, Lucknow, in November 2022, I served on a two-year extension at R.L.B. Hospital, Lucknow, which concluded in 2024. I am currently 65 years of age and am medically fit to serve as per the criteria (up to 70 years) mentioned in the notice.')

    doc.add_paragraph(
        'I have attached my Bio-data and self-attested copies of my educational and experience certificates for your kind perusal. I am willing to provide services according to the roster and telemedicine requirements as specified.')

    doc.add_paragraph(
        'I look forward to the opportunity to contribute to the urban healthcare community through this initiative.')

    doc.add_paragraph('\nYours Sincerely,')
    doc.add_paragraph('\n\n(Dr. Rakesh Singh)\nMobile: 9415136011\nAddress: D-141 Rajajipuram, Lucknow')

    doc.save('Job_Application_Dr_Rakesh_Singh.docx')


create_resume()
create_application()